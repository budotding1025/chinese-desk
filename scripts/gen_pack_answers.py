# -*- coding: utf-8 -*-
"""Generate printable A4 answer pages for 单元冲刺 + 完整摸底卷; update print-index."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from docx_utils import docx_to_pdf, set_run_font
from theme import BRAND_DEEP, BRAND_ORANGE_RED, BRAND_ORANGE_YELLOW, INK

ROOT = Path(__file__).resolve().parents[1]
OUT_SA = ROOT / "printables" / "sprint-answers"
OUT_FA = ROOT / "printables" / "full-answers"
INDEX = ROOT / "print-index.json"

SPRINT_ANSWERS = {
    "u1": {
        "title": "第一单元冲刺 · 参考答案",
        "lines": [
            "一、看拼音写词语",
            "繁星、模糊、浩浩荡荡、山崩地裂、鹅卵石、逐渐（以教材会写字词为准）",
            "二、多音 / 错读",
            "观潮 cháo；朝阳 zhāo；霎时 shà（不作 chà）；萤火虫 yíng（不作 yín）",
            "三、园地默写",
            "《鹿柴》：空山不见人，但闻人语响。返景入深林，复照青苔上。",
            "《赠刘景文》名句：一年好景君须记，最是橙黄橘绿时。（或荷尽已无擎雨盖……）",
            "四、仿写",
            "忽然……过了一会儿……（通顺、写出变化即可）",
            "五、短阅读",
            "潮来声势句：判断/选择以卷面为准；好处示例：写出潮水声势浩大、气势勇猛（意思对即可）。",
        ],
    },
    "u2": {
        "title": "第二单元冲刺 · 参考答案",
        "lines": [
            "一、字词",
            "身份、揭晓、注视、的确、证明、蝙蝠、呼风唤雨",
            "二、选字 / 多音",
            "纲/冈：提纲、井冈山、纲要；即/既：即使、既然、立即；具/俱：具体、俱乐部、俱备",
            "蝙蝠 biān；一溜烟 liù；避开 bì",
            "三、改写 + 园地",
            "改写①并列展开：……的时候……，……的时候……，……的时候也……",
            "改写②设问：是谁……呢？当然是……。",
            "园地：人非生而知之者，孰能无惑？／博学之，审问之，慎思之，明辨之，笃行之。／善疑者，不疑人之所疑，而疑人之所不疑。",
            "四、阅读",
            "解释词语意思对即可；能按内容/写法/生活分类提问即可。",
        ],
    },
    "u3": {
        "title": "第三单元冲刺 · 参考答案",
        "lines": [
            "一、字词",
            "光滑、嫩绿、爬山虎、均匀、空隙、住宅、慎重",
            "二、多音 / 错别字",
            "投降 xiáng / 降低 jiàng；歌曲 qǔ / 曲折 qū；茎 jīng；谨慎 shèn",
            "痕迹（非痕际）；平坦（非平担）；触角（非触脚）",
            "三、诗句 / 俗语",
            "横看成岭侧成峰，远近高低各不同；只缘身在此山中（缘＝因为）",
            "一场秋雨一场寒，十场秋雨要穿棉（或：十场秋雨穿上棉）",
            "四、观察仿写 + 短阅读",
            "仿写写出变化过程即可；「潜伏」：隐藏、埋伏（意思对即可）。",
        ],
    },
}

FULL_MD = {
    "u1": ROOT / "answers" / "u01-四上第一单元-参考答案.md",
    "u2": ROOT / "answers" / "u02-四上第二单元-参考答案.md",
    "u3": ROOT / "answers" / "u03-四上第三单元-参考答案.md",
}


def tight_doc():
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.4)
    st = doc.styles["Normal"]
    st.font.name = "宋体"
    st.font.size = Pt(11)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return doc


def add_p(doc, text, *, size=11, bold=False, color=None, center=False, after=3, before=0, font="宋体"):
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.12
    run = para.add_run(text)
    set_run_font(run, name_cn=font, size=size, bold=bold, color=color or INK)
    return para


def save_pair(doc, folder: Path, stem: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    docx = folder / f"{stem}.docx"
    pdf = folder / f"{stem}.pdf"
    doc.save(str(docx))
    try:
        docx_to_pdf(docx, pdf)
    except Exception as e:
        print("COM pdf fail", e)
        import subprocess

        ps = (
            f"$w=New-Object -ComObject Word.Application; $w.Visible=$false; "
            f"$d=$w.Documents.Open('{docx}'); "
            f"$d.SaveAs([string]'{pdf}',17); $d.Close(); $w.Quit()"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=False)
    print("OK", pdf.name if pdf.exists() else docx.name)
    return pdf


def build_from_lines(title: str, lines: list[str], note: str = ""):
    doc = tight_doc()
    add_p(doc, "语文书桌 · 参考答案（做完再看）", size=9, bold=True, color=BRAND_ORANGE_YELLOW, center=True, after=2)
    add_p(doc, title, size=15, bold=True, color=BRAND_ORANGE_RED, center=True, after=2, font="黑体")
    add_p(
        doc,
        "单独答案页 · 可打印 A4　|　先做卷再对答案",
        size=10,
        center=True,
        after=6,
        color=BRAND_DEEP,
    )
    if note:
        add_p(doc, note, size=9, after=6, color=BRAND_DEEP)
    for line in lines:
        raw = (line or "").strip()
        if not raw:
            continue
        if re.match(r"^#{1,3}\s+", raw):
            raw = re.sub(r"^#{1,3}\s+", "", raw)
        if re.match(r"^[一二三四五六七八九十]+[、．.]", raw) or raw.startswith("##"):
            add_p(doc, raw.lstrip("#").strip(), size=12, bold=True, color=BRAND_ORANGE_RED, before=8, after=4)
        elif raw.startswith("###"):
            add_p(doc, raw.lstrip("#").strip(), size=11, bold=True, color=BRAND_DEEP, before=6, after=3)
        elif raw.startswith(">"):
            add_p(doc, raw.lstrip("> ").strip(), size=9, color=BRAND_DEEP, after=4)
        elif raw.startswith("---"):
            continue
        else:
            # strip markdown bold markers lightly
            clean = raw.replace("**", "")
            add_p(doc, clean, size=11, after=3)
    add_p(doc, "□已对答案　□错题已入档", size=10, color=BRAND_DEEP, before=10)
    return doc


def md_to_lines(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    out = []
    for line in text.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        out.append(s)
    return out


def main():
    OUT_SA.mkdir(parents=True, exist_ok=True)
    OUT_FA.mkdir(parents=True, exist_ok=True)

    sprint_ans = {}
    for uid, meta in SPRINT_ANSWERS.items():
        stem = f"{uid.upper()}-sprint-answers"
        doc = build_from_lines(meta["title"], meta["lines"])
        pdf = save_pair(doc, OUT_SA, stem)
        sprint_ans[uid] = f"printables/sprint-answers/{stem}.pdf"
        # also write md
        md = ROOT / "answers" / f"{uid}-sprint-参考答案.md"
        md.write_text(
            f"# {meta['title']}\n\n> 对应单元冲刺卷\n\n" + "\n\n".join(meta["lines"]) + "\n",
            encoding="utf-8",
        )

    full_ans = {}
    titles = {
        "u1": "第一单元完整摸底卷 · 参考答案",
        "u2": "第二单元完整摸底卷 · 参考答案",
        "u3": "第三单元完整摸底卷 · 参考答案（第1–2页）",
    }
    stems = {
        "u1": "u01-full-answers",
        "u2": "u02-full-answers",
        "u3": "u03-full-answers",
    }
    for uid, md_path in FULL_MD.items():
        lines = md_to_lines(md_path) if md_path.exists() else ["（答案待补）"]
        # drop the top H1 duplicate; builder adds title
        if lines and lines[0].startswith("#"):
            lines = lines[1:]
        doc = build_from_lines(titles[uid], lines)
        pdf = save_pair(doc, OUT_FA, stems[uid])
        full_ans[uid] = f"printables/full-answers/{stems[uid]}.pdf"

    index = {}
    if INDEX.exists():
        index = json.loads(INDEX.read_text(encoding="utf-8"))
    index["sprintAnswers"] = sprint_ans
    index["fullAnswers"] = full_ans
    # normalize slash
    for key in ("sprint", "full", "sprintAnswers", "fullAnswers"):
        if key in index and isinstance(index[key], dict):
            index[key] = {k: str(v).replace("\\", "/") for k, v in index[key].items()}
    INDEX.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    print("index sprintAnswers", sprint_ans)
    print("index fullAnswers", full_ans)

    (OUT_SA / "README.md").write_text("# 单元冲刺 · 答案（A4）\n\n与 `../sprint/` 对应，单独打印。\n", encoding="utf-8")
    (OUT_FA / "README.md").write_text("# 完整摸底卷 · 答案（A4）\n\n与 `../full/` 对应，单独打印。\n", encoding="utf-8")


if __name__ == "__main__":
    main()
