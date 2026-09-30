# -*- coding: utf-8 -*-
"""Printable mistake log sheet + ensure mistakes folder."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx import Document
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_utils import ROOT, add_para, docx_to_pdf, header_block, new_doc, footer_line
from theme import BRAND_ORANGE_RED, INK

OUT = ROOT / "printables" / "weekday"


def build():
    doc = new_doc()
    header_block(
        doc,
        "错题登记表（可打印）",
        "做完练习马上填 → 周末誊入 mistakes/错题档案.md → 考前按单元复习",
    )
    add_para(doc, "姓名：________　　单元：U____　　来源：□工作日迷你 □周末 □整卷摸底 □错题再练", size=10)
    add_para(doc, "日期：________", size=10)

    table = doc.add_table(rows=9, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["题号", "类型", "我的错法", "正确要点", "状态", "考前再练"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
                r.font.name = "宋体"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for r in range(1, 9):
        for c in range(6):
            for p in table.rows[r].cells[c].paragraphs:
                p.paragraph_format.space_after = Pt(0)

    add_para(doc, "", space_after=6)
    add_para(doc, "类型可写：生字 / 形近组词 / 多音 / 字义 / 园地默写 / 园地理解 / 句型仿写 / 短阅读 / 作文 / 课外阅读", size=8)
    add_para(doc, "状态：未订正 / 已订正 / 已掌握", size=8)
    add_para(doc, "考试前：只复习本单元「未掌握」和勾了考前再练的题。", size=9, color=BRAND_ORANGE_RED)
    footer_line(doc, "语文书桌 · 错题登记表")
    return doc


def main():
    (ROOT / "mistakes").mkdir(exist_ok=True)
    doc = build()
    docx = OUT / "错题登记表.docx"
    pdf = OUT / "错题登记表.pdf"
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx))
    print("Wrote", docx)
    try:
        docx_to_pdf(docx, pdf)
    except Exception:
        import subprocess

        ps = (
            f"$w=New-Object -ComObject Word.Application; $w.Visible=$false; "
            f"$d=$w.Documents.Open('{docx.resolve()}'); "
            f"$d.SaveAs([ref]'{pdf.resolve()}',[ref]17); $d.Close($false); $w.Quit()"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print("Wrote", pdf)


if __name__ == "__main__":
    main()
