# -*- coding: utf-8 -*-
"""Shared Word helpers for Chinese Desk printables (橙红/橙黄主色)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from theme import BRAND_DEEP, BRAND_ORANGE_RED, BRAND_ORANGE_YELLOW, INK

ROOT = Path(__file__).resolve().parents[1]


def new_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    # A4 小学生练习：左右留白接近单元测卷
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return doc


def set_run_font(run, name_cn="宋体", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_cn
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name_cn)


def add_para(
    doc,
    text="",
    *,
    size=10.5,
    bold=False,
    space_after=4,
    space_before=0,
    align=None,
    first_line=None,
    color=None,
    font="宋体",
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    if first_line is not None:
        pf.first_line_indent = first_line
    if text:
        run = p.add_run(text)
        set_run_font(run, name_cn=font, size=size, bold=bold, color=color)
    return p


def add_runs(doc, parts, *, size=10.5, space_after=4, align=None, first_line=None):
    """parts: list of (text, bold, color) or str."""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    if first_line is not None:
        pf.first_line_indent = first_line
    for part in parts:
        if isinstance(part, str):
            text, bold, color = part, False, INK
        else:
            text, bold, color = part[0], part[1] if len(part) > 1 else False, part[2] if len(part) > 2 else INK
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def header_block(doc, title: str, note: str = ""):
    add_para(
        doc,
        "语文书桌 · Chinese Desk",
        size=9,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
        color=BRAND_ORANGE_YELLOW,
    )
    add_para(
        doc,
        title,
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
        color=BRAND_ORANGE_RED,
        font="黑体",
    )
    add_para(
        doc,
        "学校：__________　班级：__________　姓名：__________　成绩：__________",
        size=10,
        space_after=6,
        color=INK,
    )
    if note:
        add_para(doc, note, size=8, space_after=6, color=BRAND_DEEP, align=WD_ALIGN_PARAGRAPH.CENTER)


def section_title(doc, text: str):
    add_para(doc, text, size=12, bold=True, space_before=8, space_after=8, color=BRAND_ORANGE_RED)


def footer_line(doc, text: str):
    add_para(doc, text, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=0, color=BRAND_ORANGE_YELLOW)


def _set_cell_border(cell, **kwargs):
    """kwargs: top/left/bottom/right -> {'sz': '4', 'val': 'single', 'color': '000000'} or None to clear."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        spec = kwargs.get(edge)
        existing = tcBorders.find(qn(f"w:{edge}"))
        if existing is not None:
            tcBorders.remove(existing)
        if not spec:
            continue
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:sz"), str(spec.get("sz", "12")))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), spec.get("color", "000000"))
        tcBorders.append(el)


def _clear_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is not None:
        tblPr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def add_pinyin_write_grid(doc, items, *, cols=3, py_size=14, gap_after=10):
    """看拼音写词语：三列网格，拼音居中、下方留白横线一一对应（贴近单元测卷）。

    items: list of (zh, py) or py strings.
    """
    from docx.enum.table import WD_TABLE_ALIGNMENT

    pairs = []
    for it in items:
        if isinstance(it, (tuple, list)):
            zh, py = it[0], it[1]
        else:
            zh, py = "", str(it)
        pairs.append((zh, py))

    # usable width ≈ A4 21cm − margins 1.8*2 ≈ 17.4cm
    usable = 17.2
    col_w = usable / cols

    for start in range(0, len(pairs), cols):
        chunk = pairs[start : start + cols]
        table = doc.add_table(rows=1, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _clear_table_borders(table)

        for i in range(cols):
            cell = table.cell(0, i)
            cell.width = Cm(col_w)
            cell.text = ""
            para_py = cell.paragraphs[0]
            para_py.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para_py.paragraph_format.space_before = Pt(2)
            para_py.paragraph_format.space_after = Pt(6)
            para_py.paragraph_format.line_spacing = 1.0

            if i < len(chunk):
                zh, py = chunk[i]
                run = para_py.add_run(py)
                set_run_font(run, name_cn="宋体", size=py_size, bold=False, color=INK)
                rPr = run._element.get_or_add_rPr()
                rFonts = rPr.get_or_add_rFonts()
                rFonts.set(qn("w:ascii"), "Times New Roman")
                rFonts.set(qn("w:hAnsi"), "Times New Roman")

                n = max(2, len(zh) if zh else 2)
                para_line = cell.add_paragraph()
                para_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para_line.paragraph_format.space_before = Pt(0)
                para_line.paragraph_format.space_after = Pt(4)
                pPr = para_line._p.get_or_add_pPr()
                pBdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "18")
                bottom.set(qn("w:space"), "10")
                bottom.set(qn("w:color"), "000000")
                pBdr.append(bottom)
                pPr.append(pBdr)
                # 按字数加宽留白，方便小学生书写
                spacer = para_line.add_run("　" * (n + 3))
                set_run_font(spacer, size=16, color=INK)

            _set_cell_border(cell, top=None, left=None, bottom=None, right=None)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(gap_after)


def exam_doc() -> Document:
    """A4 完整摸底卷：边距贴近小学单元测原卷。"""
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return doc


def exam_header(doc, title: str, *, incomplete: bool = False):
    """原卷式页眉：标题 + 学校班级姓名成绩（少装饰、少占高）。"""
    add_para(
        doc,
        title,
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
        space_before=0,
        color=INK,
        font="黑体",
    )
    add_para(
        doc,
        "学校：____________　班级：________　姓名：____________　成绩：________",
        size=10.5,
        space_after=8 if not incomplete else 3,
        color=INK,
    )
    if incomplete:
        add_para(
            doc,
            "（扫描件仅有第 1–2 页 · 第 3–4 页待补）",
            size=8,
            space_after=6,
            color=BRAND_DEEP,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )


def exam_section(doc, text: str):
    add_para(doc, text, size=11, bold=True, space_before=8, space_after=4, color=INK)


def exam_body(doc, text: str, *, size=10.5, space_after=2, space_before=0, first_line=None, bold=False):
    return add_para(
        doc,
        text,
        size=size,
        bold=bold,
        space_after=space_after,
        space_before=space_before,
        first_line=first_line,
        color=INK,
    )


def exam_footer(doc, text: str):
    add_para(doc, text, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=0, color=INK)


def write_lines(doc, n: int = 1, *, size=12, space_after=2):
    """作答横线：字号略大，方便手写。"""
    for _ in range(n):
        add_para(
            doc,
            "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿",
            size=size,
            space_after=space_after,
            color=INK,
        )


def _cell_v_center(cell):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_boxes_in_cell(cell, n: int, *, box_cm: float = 0.7):
    """在单元格内画 n 个田字格（外框实线、内十字虚线），贴近原卷。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH as ALG

    outer = {"sz": "12", "val": "single", "color": "000000"}
    dash = {"sz": "6", "val": "dashed", "color": "808080"}
    host = cell.add_table(rows=1, cols=n)
    host.autofit = False
    for i in range(n):
        outer_cell = host.cell(0, i)
        outer_cell.width = Cm(box_cm)
        outer_cell.text = ""
        # clear default paragraph height
        op = outer_cell.paragraphs[0]
        op.paragraph_format.space_before = Pt(0)
        op.paragraph_format.space_after = Pt(0)
        # 2x2 田字格
        grid = outer_cell.add_table(rows=2, cols=2)
        grid.autofit = False
        half = box_cm / 2
        for rr in range(2):
            for cc in range(2):
                gc = grid.cell(rr, cc)
                gc.width = Cm(half)
                gc.text = ""
                gp = gc.paragraphs[0]
                gp.paragraph_format.space_before = Pt(0)
                gp.paragraph_format.space_after = Pt(0)
                gr = gp.add_run(" ")
                set_run_font(gr, size=6, color=INK)
                # outer edges solid, inner edges dashed
                top = outer if rr == 0 else dash
                bottom = outer if rr == 1 else dash
                left = outer if cc == 0 else dash
                right = outer if cc == 1 else dash
                _set_cell_border(gc, top=top, left=left, bottom=bottom, right=right)
        # exact row heights for 田字格
        for row in grid.rows:
            tr = row._tr
            trPr = tr.get_or_add_trPr()
            trHeight = OxmlElement("w:trHeight")
            trHeight.set(qn("w:val"), str(int(half * 567)))
            trHeight.set(qn("w:hRule"), "exact")
            trPr.append(trHeight)
        _set_cell_border(outer_cell, top=None, left=None, bottom=None, right=None)
    tr = host.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(box_cm * 567)))
    trHeight.set(qn("w:hRule"), "exact")
    trPr.append(trHeight)


def add_inline_pinyin_line(doc, segments, *, size=10.5, box_cm: float = 0.7, after: float = 3):
    """一行内嵌拼音方格（与 JPG 原卷一致：拼音在格上、格嵌在句中）。

    segments: str 或 (pinyin, n_chars)
    """
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH as ALG

    specs = []
    for seg in segments:
        if isinstance(seg, tuple):
            specs.append(("blank", seg[0], int(seg[1])))
        elif seg:
            specs.append(("text", str(seg), 0))
    if not specs:
        return

    table = doc.add_table(rows=1, cols=len(specs))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    _clear_table_borders(table)

    for i, spec in enumerate(specs):
        cell = table.cell(0, i)
        _cell_v_center(cell)
        _set_cell_border(cell, top=None, left=None, bottom=None, right=None)
        cell.text = ""
        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)

        if spec[0] == "text":
            # 勿压窄：宁可略宽，避免汉字被裁切
            text = spec[1]
            cell.width = Cm(max(0.8, min(14.0, 0.52 * max(1, len(text)) + 0.3)))
            p0.alignment = ALG.LEFT
            run = p0.add_run(text)
            set_run_font(run, size=size, color=INK)
        else:
            py, n = spec[1], spec[2]
            cell.width = Cm(box_cm * n + 0.2)
            p0.alignment = ALG.CENTER
            p0.paragraph_format.space_after = Pt(0)
            run = p0.add_run(py)
            set_run_font(run, size=8.5, color=INK)
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn("w:ascii"), "Times New Roman")
            rFonts.set(qn("w:hAnsi"), "Times New Roman")
            _add_boxes_in_cell(cell, n, box_cm=box_cm)

    # row height: pinyin + box, 贴近原卷行高
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int((box_cm + 0.38) * 567)))
    trHeight.set(qn("w:hRule"), "atLeast")
    trPr.append(trHeight)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(after)


def add_pinyin_sentence(doc, segments, *, size=10.5):
    """兼容旧调用：整段按一行内嵌格排（长句请在生成端拆成多行调用 add_inline_pinyin_line）。"""
    add_inline_pinyin_line(doc, segments, size=size, box_cm=0.7, after=4)


def add_choice_char_table(doc, pairs_and_blanks):
    """选字组词：左侧字对竖栏 + 右侧三列填空（贴近 U2 原卷）。"""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH as ALG

    rows = len(pairs_and_blanks)
    table = doc.add_table(rows=rows, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    border = {"sz": "10", "val": "single", "color": "000000"}
    # merge left column into one box
    if rows > 1:
        table.cell(0, 0).merge(table.cell(rows - 1, 0))
    c0 = table.cell(0, 0)
    c0.width = Cm(2.4)
    c0.text = ""
    for idx, (pair, _) in enumerate(pairs_and_blanks):
        p = c0.paragraphs[0] if idx == 0 else c0.add_paragraph()
        p.alignment = ALG.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(pair)
        set_run_font(run, size=10.5, color=INK)
    _set_cell_border(c0, top=border, left=border, bottom=border, right=border)

    for r, (_, blanks) in enumerate(pairs_and_blanks):
        for j in range(3):
            c = table.cell(r, j + 1)
            c.width = Cm(4.7)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            txt = blanks[j] if j < len(blanks) else ""
            run = p.add_run(txt)
            set_run_font(run, size=10.5, color=INK)
            _set_cell_border(c, top=None, left=None, bottom=None, right=None)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(6)


def add_emphasis_words(doc, label: str, words: list[str], *, size=10.5):
    """加点词语：Word 着重号。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(label)
    set_run_font(run, size=size, color=INK)
    for w in words:
        run = p.add_run(w)
        set_run_font(run, size=size, color=INK)
        rPr = run._element.get_or_add_rPr()
        em = OxmlElement("w:em")
        em.set(qn("w:val"), "dot")
        rPr.append(em)
        run2 = p.add_run("　")
        set_run_font(run2, size=size, color=INK)


def add_page_break(doc):
    doc.add_page_break()


def docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx_path.resolve()))
        d.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        d.Close(False)
    finally:
        word.Quit()
