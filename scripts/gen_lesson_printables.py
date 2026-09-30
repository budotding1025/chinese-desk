# -*- coding: utf-8 -*-
"""Per-lesson A4 printables (~10 min) + recite sheet (+10) + unit sprint."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from docx_utils import add_pinyin_write_grid, docx_to_pdf, set_run_font
from theme import BRAND_DEEP, BRAND_ORANGE_RED, BRAND_ORANGE_YELLOW, INK

ROOT = Path(__file__).resolve().parents[1]
OUT_L = ROOT / "printables" / "lessons"
OUT_S = ROOT / "printables" / "sprint"
ANS = ROOT / "answers"

# Keep printable content in sync with app data.js (authoritative for print generation)
LESSONS = [
    {
        "unit": "u1",
        "unitName": "第一单元",
        "bookNo": 1,
        "title": "观潮",
        "words": [("潮汐", "cháo xī"), ("据说", "jù shuō"), ("大堤", "dà dī"), ("盼望", "pàn wàng"), ("逐渐", "zhú jiàn"), ("霎时", "shà shí")],
        "idioms": [("浩浩荡荡", "hào hào dàng dàng"), ("山崩地裂", "shān bēng dì liè"), ("齐头并进", "qí tóu bìng jìn")],
        "compounds": [("潮", "朝"), ("堤", "提"), ("盼", "扮")],
        "poly": [("观潮", ["cháo", "zhāo"], 0), ("朝阳", ["cháo", "zhāo"], 1), ("闷雷", ["mēn", "mèn"], 1), ("霎时", ["shà", "chà"], 0)],
        "poly_wrong": [("薄雾", "bó", True), ("萤火虫", "yín", False), ("芦苇", "wěi", True)],  # pick wrong-reading style
        "pattern": {
            "label": "用上「霎时／顿时／过了一会儿」写潮来或下雨（2～3 句）",
            "example": "例：远处出现一条白线，霎时拉长变粗；过了一会儿，浪潮齐头并进。",
        },
        "recite": {
            "label": "背诵第3～4自然段 · 默写（另加约10分钟）",
            "lines": [
                "浪潮越来越近，犹如千万匹白色战马____________________，",
                "____________________地飞奔而来；那声音如同____________________，",
                "好像大地都被震得颤动起来。",
            ],
            "answer": "齐头并进；浩浩荡荡；山崩地裂",
        },
    },
    {
        "unit": "u1",
        "unitName": "第一单元",
        "bookNo": 2,
        "title": "走月亮",
        "words": [("鹅卵石", "é luǎn shí"), ("坑坑洼洼", "kēng keng wā wā"), ("稻穗", "dào suì"), ("成熟", "chéng shú"), ("葡萄", "pú tao"), ("闪烁", "shǎn shuò")],
        "idioms": [],
        "compounds": [("卵", "卯"), ("稻", "蹈"), ("填", "慎")],
        "poly": [("成熟", ["shú", "shóu"], 0)],
        "poly_wrong": [],
        "pattern": {
            "label": "仿写：每个小水塘都抱着一个月亮！（用拟人写月下景物一句）",
            "example": "例：每扇窗户都盛着一轮月亮。",
        },
        "recite": {
            "label": "背诵第4自然段 · 默写（另加约10分钟）",
            "lines": [
                "哟，卵石间有多少可爱的小水塘啊，每个小水塘都抱着一个____________________！",
                "看，稻谷就要成熟了，稻穗____________________，稻田像一块月光镀亮的____________________。",
            ],
            "answer": "月亮；低垂着头；银毯",
        },
    },
    {
        "unit": "u1",
        "unitName": "第一单元",
        "bookNo": 3,
        "title": "现代诗二首",
        "words": [("秋晚", "qiū wǎn"), ("归巢", "guī cháo"), ("花牛", "huā niú"), ("草地", "cǎo dì")],
        "idioms": [],
        "compounds": [("巢", "剿")],
        "poly": [],
        "poly_wrong": [],
        "pattern": {
            "label": "仿照诗句，用一个比喻写傍晚或小动物（1～2 句）",
            "example": "例：晚霞像一匹散开的红绸，轻轻盖在河面上。",
        },
        "recite": None,
    },
    {
        "unit": "u1",
        "unitName": "第一单元",
        "bookNo": 4,
        "title": "繁星",
        "words": [("繁星", "fán xīng"), ("模糊", "mó hu"), ("飞舞", "fēi wǔ"), ("柔和", "róu hé"), ("梦幻", "mèng huàn")],
        "idioms": [("摇摇欲坠", "yáo yáo yù zhuì"), ("半明半昧", "bàn míng bàn mèi"), ("密密麻麻", "mì mì má má")],
        "compounds": [("昧", "妹"), ("坠", "堕")],
        "poly": [("模糊", ["mó", "mú"], 0)],
        "poly_wrong": [("萤火虫", "yíng", True), ("半明半昧", "mèi", True), ("霎时", "chà", False)],
        "pattern": {
            "label": "用上「渐渐地……我好像……」写夜空或雨景一句",
            "example": "例：渐渐地我的眼睛模糊了，我好像看见无数萤火虫在飞舞。",
        },
        "recite": None,
    },
    {
        "unit": "u2",
        "unitName": "第二单元",
        "bookNo": 5,
        "title": "一个豆荚里的五粒豆",
        "words": [("豌豆", "wān dòu"), ("舒适", "shū shì"), ("僵硬", "jiāng yìng"), ("囚犯", "qiú fàn"), ("揭晓", "jiē xiǎo"), ("青苔", "qīng tái")],
        "idioms": [],
        "compounds": [("豌", "碗"), ("僵", "疆"), ("溢", "隘")],
        "poly": [("囚犯", ["qiú", "qiū"], 0)],
        "poly_wrong": [],
        "pattern": {
            "label": "照样子提问：针对内容 / 写法各提一个问题",
            "example": "例：内容——小女孩为什么慢慢好起来？写法——为什么把豌豆比作囚犯？",
        },
        "recite": None,
    },
    {
        "unit": "u2",
        "unitName": "第二单元",
        "bookNo": 6,
        "title": "夜间飞行的秘密",
        "words": [("蝙蝠", "biān fú"), ("超声波", "chāo shēng bō"), ("障碍", "zhàng ài"), ("敏锐", "mǐn ruì"), ("绳子", "shéng zi")],
        "idioms": [("一溜烟", "yí liù yān")],
        "compounds": [("蝠", "福"), ("绳", "蝇")],
        "poly": [("蝙蝠", ["biān", "biǎn"], 0), ("一溜烟", ["liū", "liù"], 1), ("避开", ["bì", "pì"], 0)],
        "poly_wrong": [("荧屏", "yíng", True), ("嚷嚷", "rāng", True)],
        "pattern": {
            "label": "照样子改写：那盏灯晴天、阴天、雨天的时候都亮着。→ 分开说三遍「……的时候……」",
            "example": "原句：雷达白天、夜晚、雾天都能工作。改写：________________",
        },
        "recite": None,
    },
    {
        "unit": "u2",
        "unitName": "第二单元",
        "bookNo": 7,
        "title": "呼风唤雨的世纪",
        "words": [("世纪", "shì jì"), ("奥秘", "ào mì"), ("技术", "jì shù"), ("改变", "gǎi biàn")],
        "idioms": [("呼风唤雨", "hū fēng huàn yǔ")],
        "compounds": [("唤", "换"), ("纪", "记")],
        "poly": [],
        "poly_wrong": [],
        "pattern": {
            "label": "照样子：人类呼风唤雨。→ 是谁呼风唤雨呢？当然是人类。",
            "example": "原句：科学家揭开了自然的奥秘。改写：________________",
        },
        "recite": None,
    },
    {
        "unit": "u2",
        "unitName": "第二单元",
        "bookNo": 8,
        "title": "蝴蝶的家",
        "words": [("蝴蝶", "hú dié"), ("避雨", "bì yǔ"), ("珍惜", "zhēn xī"), ("安然", "ān rán"), ("忧愁", "yōu chóu")],
        "idioms": [],
        "compounds": [("蝶", "碟"), ("忧", "优")],
        "poly": [("避开", ["bì", "pì"], 0)],
        "poly_wrong": [],
        "pattern": {
            "label": "仿写：为小动物着急的两三句话（用上「怎么……呢？」）",
            "example": "例：下这么大的雨，蚂蚁的家怎么会不被冲走呢？",
        },
        "recite": None,
    },
    {
        "unit": "u3",
        "unitName": "第三单元",
        "bookNo": 9,
        "title": "古诗三首",
        "words": [("庐山", "lú shān"), ("真珠", "zhēn zhū"), ("逊色", "xùn sè")],
        "idioms": [],
        "compounds": [("逊", "孙")],
        "poly": [],
        "poly_wrong": [],
        "pattern": {
            "label": "用自己的话解释「只缘身在此山中」，并仿一句「因为……所以看不清……」",
            "example": "例：因为就在事情中间，所以看不清全部。",
        },
        "recite": {
            "label": "默写《题西林壁》名句（另加约10分钟）",
            "lines": [
                "横看成岭侧成峰，____________________。",
                "不识庐山真面目，____________________。",
            ],
            "answer": "远近高低各不同；只缘身在此山中",
        },
    },
    {
        "unit": "u3",
        "unitName": "第三单元",
        "bookNo": 10,
        "title": "爬山虎的脚",
        "words": [("光滑", "guāng huá"), ("嫩绿", "nèn lǜ"), ("爬山虎", "pá shān hǔ"), ("均匀", "jūn yún"), ("空隙", "kòng xì"), ("叶柄", "yè bǐng")],
        "idioms": [],
        "compounds": [("均", "钧"), ("隙", "狭")],
        "poly": [("曲折", ["qǔ", "qū"], 1), ("空隙", ["kòng", "kōng"], 0), ("投降", ["jiàng", "xiáng"], 1), ("降低", ["jiàng", "xiáng"], 0)],
        "poly_wrong": [("痕迹", "jì", True), ("触角", "jiǎo", True)],
        "pattern": {
            "label": "观察仿写：像「细丝」「小逗号」那样，用比喻写一种植物的一部分",
            "example": "例：新长的叶子像一把把撑开的小绿伞。",
        },
        "recite": None,
    },
    {
        "unit": "u3",
        "unitName": "第三单元",
        "bookNo": 11,
        "title": "蟋蟀的住宅",
        "words": [("住宅", "zhù zhái"), ("洞穴", "dòng xué"), ("慎重", "shèn zhòng"), ("柔弱", "róu ruò"), ("搜索", "sōu suǒ")],
        "idioms": [],
        "compounds": [("宅", "诧"), ("慎", "填")],
        "poly": [("慎重", ["shèn", "chén"], 0)],
        "poly_wrong": [],
        "pattern": {
            "label": "仿写：先写外形/位置，再写一个具体活动（观察小动物）",
            "example": "例：它的洞穴朝阳，出口不积水；它用柔弱的腿一下一下扒土。",
        },
        "recite": None,
    },
]

SPRINTS = [
    {
        "unit": "u1",
        "title": "第一单元冲刺",
        "minutes": "20–25",
        "blocks": [
            "一、看拼音写词语（本单元）：繁星、模糊、浩浩荡荡、山崩地裂、鹅卵石、逐渐",
            "二、多音/错读：观潮（cháo）朝阳（zhāo）；霎时不作 chà；萤火虫不作 yín",
            "三、园地：默写《鹿柴》或《赠刘景文》名句",
            "四、仿写：忽然……过了一会儿……",
            "五、短阅读：潮来声势句 → 判断/选择/一句说好处",
        ],
    },
    {
        "unit": "u2",
        "title": "第二单元冲刺",
        "minutes": "20–25",
        "blocks": [
            "一、字词：身份、揭晓、注视、的确、证明、蝙蝠、呼风唤雨",
            "二、选字：纲/冈、即/既、具/俱；多音：蝙蝠 biān、一溜烟 liù、避开 bì",
            "三、改写句子两类；园地名句默写（韩愈、中庸、方以智）",
            "四、阅读：解释词语 + 分类提问",
        ],
    },
    {
        "unit": "u3",
        "title": "第三单元冲刺",
        "minutes": "20–25",
        "blocks": [
            "一、字词：光滑、嫩绿、爬山虎、均匀、空隙、住宅、慎重",
            "二、多音：降 xiáng/jiàng、曲 qǔ/qū、茎 jīng；错别字：痕迹、平坦、触角",
            "三、诗句：题西林壁、秋雨俗语；「缘」=因为",
            "四、观察仿写 + 短阅读「潜伏」",
        ],
    },
]


def tight_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = "宋体"
    st.font.size = Pt(12)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return doc


def p(doc, text, *, size=12, bold=False, color=None, center=False, after=3, before=0):
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.2
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color or INK)
    return para


def section(doc, t):
    p(doc, t, size=12, bold=True, color=BRAND_ORANGE_RED, before=8, after=8)


def blank(doc, n=1):
    for _ in range(n):
        p(doc, "＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿", size=14, after=10)


def save_pair(doc, folder: Path, stem: str):
    folder.mkdir(parents=True, exist_ok=True)
    docx = folder / f"{stem}.docx"
    pdf = folder / f"{stem}.pdf"
    doc.save(str(docx))
    try:
        docx_to_pdf(docx, pdf)
    except Exception:
        import subprocess

        ps = (
            f"$w=New-Object -ComObject Word.Application; $w.Visible=$false; "
            f"$d=$w.Documents.Open('{docx.resolve()}'); "
            f"$d.SaveAs([ref]'{pdf.resolve()}',[ref]17); $d.Close($false); $w.Quit()"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print("OK", stem)


def build_lesson(les: dict) -> Document:
    doc = tight_doc()
    p(doc, "语文书桌 · Chinese Desk", size=10, bold=True, color=BRAND_ORANGE_YELLOW, center=True, after=2)
    p(
        doc,
        f"第{les['bookNo']}课《{les['title']}》· 每日练习（约10–15分钟）",
        size=16,
        bold=True,
        color=BRAND_ORANGE_RED,
        center=True,
        after=2,
    )
    p(doc, f"{les['unitName']}　|　A4　|　姓名：________　日期：________　用时：____分钟", size=11, center=True, after=8)

    # 控制在约10–15分钟：字词≤9、组词≤3、多音≤3、仿写1句
    words = (les.get("words") or [])[:9]
    compounds = (les.get("compounds") or [])[:3]
    polys = (les.get("poly") or [])[:3]

    section(doc, "一、字词听写（看拼音写，最多9个）")
    add_pinyin_write_grid(doc, words, cols=3, py_size=14, gap_after=10)

    if les.get("idioms"):
        section(doc, "二、成语（看拼音写）")
        add_pinyin_write_grid(doc, les["idioms"][:4], cols=2, py_size=14, gap_after=10)
    elif compounds:
        section(doc, "二、形近组词（最多3组）")
        for a, b in compounds:
            p(doc, f"{a}（　　　　）　　{b}（　　　　）", size=14, after=8)
    else:
        section(doc, "二、组词")
        p(doc, "（本组略）", size=11, after=4)

    section(doc, "三、多音字（最多3个，正确读音下画√）")
    if polys:
        for word, opts, _ in polys:
            p(doc, f"{word}　（{'　'.join(opts)}）", size=14, after=6)
    elif les.get("poly_wrong"):
        bits = [f"{i}.{w}（{py}）" for i, (w, py, ok) in enumerate(les["poly_wrong"][:3], 1)]
        p(doc, "　".join(bits), size=12, after=4)
        p(doc, "有误：第____项，应读________", size=12, after=6)
    else:
        p(doc, "（本课多音较少，本项可跳过）", size=11, after=4)

    section(doc, "四、句型仿写（写1句即可）")
    pat = les["pattern"]
    p(doc, pat["label"], size=12, after=3)
    p(doc, pat["example"], size=11, color=BRAND_DEEP, after=4)
    blank(doc, 1)

    p(doc, "□错题记入错题档案　　□已订正", size=10, color=BRAND_DEEP, before=6)
    p(doc, f"第{les['bookNo']}课 · 约10–15分钟 · 有背诵另做默写单", size=9, color=BRAND_ORANGE_YELLOW, center=True, before=6)
    return doc


def build_recite(les: dict) -> Document:
    r = les["recite"]
    doc = tight_doc()
    p(doc, "语文书桌 · Chinese Desk", size=9, bold=True, color=BRAND_ORANGE_YELLOW, center=True, after=1)
    p(
        doc,
        f"第{les['bookNo']}课《{les['title']}》· 课文默写（约10–15分钟）",
        size=14,
        bold=True,
        color=BRAND_ORANGE_RED,
        center=True,
        after=1,
    )
    p(doc, f"{r['label']}　|　姓名：________　日期：________", size=9, center=True, after=6)
    section(doc, "一、根据课文填空（不看书 · 约10分钟）")
    for line in r["lines"][:6]:
        p(doc, line, size=11, after=5)
    blank(doc, 1)
    p(doc, "（整段默写改日再练，今日只做填空即可）", size=10, color=BRAND_DEEP, after=4)
    p(doc, "□背诵过关　　□错字已订正", size=9, color=BRAND_DEEP, before=6)
    p(doc, "参考答案见 answers/lesson-printables-参考答案.md（做完再看）", size=8, color=BRAND_ORANGE_YELLOW, center=True)
    return doc


def build_sprint(sp: dict) -> Document:
    doc = tight_doc()
    p(doc, "语文书桌 · Chinese Desk", size=9, bold=True, color=BRAND_ORANGE_YELLOW, center=True, after=1)
    p(doc, f"{sp['title']}（可打印 · 约{sp['minutes']}分钟）", size=14, bold=True, color=BRAND_ORANGE_RED, center=True, after=1)
    p(doc, "学校：________　班级：________　姓名：________　成绩：________", size=10, after=4)
    p(doc, "测前冲刺：先做本卷 → 对答案 → 错题入档 → 只补漏洞。完整摸底卷见 printables/full/。", size=8, color=BRAND_DEEP, after=6)
    for i, block in enumerate(sp["blocks"], 1):
        section(doc, block if block.startswith(("一", "二", "三", "四", "五")) else f"{i}、{block}")
        # expand into writable area
        if "看拼音" in block or "字词" in block:
            p(doc, "请默写本单元易错词语（每词写一遍）：", size=10, after=3)
            blank(doc, 3)
        elif "多音" in block or "选字" in block:
            p(doc, "在正确读音下画√ / 选字填空（可另纸）：", size=10, after=3)
            blank(doc, 2)
        elif "园地" in block or "诗句" in block or "名句" in block:
            p(doc, "默写区：", size=10, after=3)
            blank(doc, 3)
        elif "仿写" in block or "改写" in block or "观察" in block:
            blank(doc, 3)
        else:
            blank(doc, 2)
    p(doc, "□错题已入档　□园地已过关", size=9, color=BRAND_DEEP, before=4)
    return doc


def write_answers():
    lines = ["# 每课打印练习 · 参考答案\n\n> 做完再看。课文原句以教材为准。\n"]
    for les in LESSONS:
        lines.append(f"\n## 第{les['bookNo']}课《{les['title']}》\n")
        lines.append("- 听写：" + "、".join(w for w, _ in les["words"]) + "\n")
        if les.get("idioms"):
            lines.append("- 成语：" + "、".join(w for w, _ in les["idioms"]) + "\n")
        if les.get("compounds"):
            lines.append("- 组词：略（合理即可）\n")
        if les.get("poly"):
            ans = "；".join(f"{w}-{opts[i]}" for w, opts, i in les["poly"])
            lines.append(f"- 多音：{ans}\n")
        if les.get("recite"):
            lines.append(f"- 默写：{les['recite']['answer']}\n")
    (ANS / "lesson-printables-参考答案.md").write_text("".join(lines), encoding="utf-8")


def patch_data_js_print_paths():
    """Inject printables paths into print-index.json (preserve answer fields if present)."""
    path = ROOT / "print-index.json"
    existing = {}
    if path.exists():
        try:
            existing = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            existing = {}
    old_lessons = (existing.get("lessons") or {}) if isinstance(existing, dict) else {}
    idx = {
        "lessons": {},
        "sprint": {
            "u1": "printables/sprint/U1-第一单元冲刺.pdf",
            "u2": "printables/sprint/U2-第二单元冲刺.pdf",
            "u3": "printables/sprint/U3-第三单元冲刺.pdf",
        },
    }
    for les in LESSONS:
        n = les["bookNo"]
        stem = f"L{n:02d}-{les['title']}-每日10分钟"
        entry = {
            "daily": f"printables/lessons/{stem}.pdf",
            "recite": f"printables/lessons/L{n:02d}-{les['title']}-背诵默写10分钟.pdf" if les.get("recite") else None,
            "title": les["title"],
        }
        prev = old_lessons.get(str(n)) or {}
        for k in ("answerPdf", "answerMd", "answerHtml"):
            if prev.get(k):
                entry[k] = prev[k]
        # default answer paths even if answers not yet regenerated
        if not entry.get("answerPdf"):
            entry["answerPdf"] = f"printables/answers/L{n:02d}-{les['title']}-答案.pdf"
        if not entry.get("answerMd"):
            entry["answerMd"] = f"answers/lessons/L{n:02d}-{les['title']}-答案.md"
        idx["lessons"][str(n)] = entry
    path.write_text(json.dumps(idx, ensure_ascii=False, indent=2), encoding="utf-8")
    print("wrote print-index.json")


def main():
    OUT_L.mkdir(parents=True, exist_ok=True)
    OUT_S.mkdir(parents=True, exist_ok=True)
    for les in LESSONS:
        stem = f"L{les['bookNo']:02d}-{les['title']}-每日10分钟"
        save_pair(build_lesson(les), OUT_L, stem)
        if les.get("recite"):
            save_pair(build_recite(les), OUT_L, f"L{les['bookNo']:02d}-{les['title']}-背诵默写10分钟")
    for sp in SPRINTS:
        save_pair(build_sprint(sp), OUT_S, f"{sp['unit'].upper()}-{sp['title']}")
    write_answers()
    patch_data_js_print_paths()
    # readme
    (OUT_L / "README.md").write_text(
        "# 每课打印（A4）\n\n- `Lxx-课题-每日10分钟`：听写、组词/成语、多音、仿写\n- `Lxx-课题-背诵默写10分钟`：有背诵要求才有\n",
        encoding="utf-8",
    )
    (OUT_S / "README.md").write_text("# 单元冲刺（A4）\n\n测前打印；完整摸底卷仍见 `../full/`。\n", encoding="utf-8")
    print("done")


if __name__ == "__main__":
    main()
